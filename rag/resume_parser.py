"""
Resume and document parser for candidate context.
Extracts text and identifies common section structures from PDF, DOCX, and TXT files.
Failure-isolated: corrupt or unsupported files raise controlled ResumeParseError.
"""
from __future__ import annotations
import io
import os
import re
from pathlib import Path
from typing import Dict, Optional
import fitz  # PyMuPDF
import docx

from rag.models import ParsedDocument


class ResumeParseError(Exception):
    """Raised when a candidate document cannot be parsed or is corrupted."""
    pass


class ResumeParser:
    """
    Parses candidate resumes and background documents.
    """

    SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}

    @classmethod
    def parse_file(cls, file_path: str | Path) -> ParsedDocument:
        path = Path(file_path)
        ext = path.suffix.lower()
        if ext not in cls.SUPPORTED_EXTENSIONS:
            raise ResumeParseError(
                f"Unsupported document format '{ext}'. Supported formats: {sorted(list(cls.SUPPORTED_EXTENSIONS))}"
            )

        if not path.exists():
            raise ResumeParseError(f"Document file does not exist: {file_path}")

        try:
            with open(path, "rb") as f:
                content = f.read()
            return cls.parse_bytes(content, filename=path.name)
        except ResumeParseError:
            raise
        except Exception as exc:
            raise ResumeParseError(f"Failed to read document '{path.name}': {type(exc).__name__}") from exc

    @classmethod
    def parse_bytes(cls, content: bytes, filename: str) -> ParsedDocument:
        if not content or len(content.strip()) == 0:
            raise ResumeParseError(f"Document '{filename}' is empty.")

        ext = Path(filename).suffix.lower()
        if ext == ".pdf":
            return cls._parse_pdf(content, filename)
        elif ext == ".docx":
            return cls._parse_docx(content, filename)
        elif ext in {".txt", ".md"}:
            return cls._parse_txt(content, filename)
        else:
            raise ResumeParseError(f"Unsupported document extension '{ext}' for file '{filename}'.")

    @classmethod
    def parse_text(cls, text: str, filename: str = "candidate_resume.txt") -> ParsedDocument:
        if not text or not text.strip():
            raise ResumeParseError("Resume text is empty.")
        clean_text = cls._normalize_text(text)
        sections = cls._extract_sections(clean_text)
        return ParsedDocument(
            filename=filename,
            format="txt",
            raw_text=clean_text,
            sections=sections,
            metadata={"character_count": len(clean_text), "word_count": len(clean_text.split())},
        )

    @classmethod
    def _parse_txt(cls, content: bytes, filename: str) -> ParsedDocument:
        try:
            text = content.decode("utf-8")
        except UnicodeDecodeError:
            try:
                text = content.decode("latin-1")
            except Exception as exc:
                raise ResumeParseError(f"Could not decode text file '{filename}': {exc}") from exc

        return cls.parse_text(text, filename=filename)

    @classmethod
    def _parse_pdf(cls, content: bytes, filename: str) -> ParsedDocument:
        try:
            doc = fitz.open(stream=content, filetype="pdf")
            pages_text = []
            for page in doc:
                pages_text.append(page.get_text())
            doc.close()
            full_text = "\n\n".join(pages_text)
        except Exception as exc:
            raise ResumeParseError(f"Corrupt or invalid PDF document '{filename}': {type(exc).__name__}") from exc

        if not full_text.strip():
            raise ResumeParseError(f"PDF document '{filename}' contains no extractable text.")

        clean_text = cls._normalize_text(full_text)
        sections = cls._extract_sections(clean_text)
        return ParsedDocument(
            filename=filename,
            format="pdf",
            raw_text=clean_text,
            sections=sections,
            metadata={"character_count": len(clean_text), "word_count": len(clean_text.split())},
        )

    @classmethod
    def _parse_docx(cls, content: bytes, filename: str) -> ParsedDocument:
        try:
            doc_file = io.BytesIO(content)
            doc = docx.Document(doc_file)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            full_text = "\n\n".join(paragraphs)
        except Exception as exc:
            raise ResumeParseError(f"Corrupt or invalid DOCX document '{filename}': {type(exc).__name__}") from exc

        if not full_text.strip():
            raise ResumeParseError(f"DOCX document '{filename}' contains no extractable text.")

        clean_text = cls._normalize_text(full_text)
        sections = cls._extract_sections(clean_text)
        return ParsedDocument(
            filename=filename,
            format="docx",
            raw_text=clean_text,
            sections=sections,
            metadata={"character_count": len(clean_text), "word_count": len(clean_text.split())},
        )

    @staticmethod
    def _normalize_text(text: str) -> str:
        # Normalize whitespace and clean null characters
        cleaned = re.sub(r"\x00", "", text)
        cleaned = re.sub(r"\r\n", "\n", cleaned)
        cleaned = re.sub(r"[ \t]+", " ", cleaned)
        cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
        return cleaned.strip()

    @staticmethod
    def _extract_sections(text: str) -> Dict[str, str]:
        """
        Heuristically breaks resume text into major sections (e.g. experience, education, skills, projects).
        """
        section_headers = [
            "summary", "profile", "about",
            "experience", "work experience", "employment",
            "skills", "technical skills", "technologies",
            "projects", "personal projects", "key projects",
            "education", "academic background",
            "certifications", "achievements",
        ]
        pattern = r"(?im)^(?:\s*)(?:" + "|".join(re.escape(h) for h in section_headers) + r")\s*[:\n]"

        matches = list(re.finditer(pattern, text))
        if not matches:
            return {"general": text}

        sections: Dict[str, str] = {}
        # Pre-header content
        if matches[0].start() > 0:
            pre_text = text[:matches[0].start()].strip()
            if pre_text:
                sections["header"] = pre_text

        for i, match in enumerate(matches):
            raw_header = match.group(0).strip().rstrip(":").lower()
            start_pos = match.end()
            end_pos = matches[i + 1].start() if i + 1 < len(matches) else len(text)
            section_content = text[start_pos:end_pos].strip()
            if section_content:
                sections[raw_header] = section_content

        return sections
